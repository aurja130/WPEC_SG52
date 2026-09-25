#!/usr/bin/env python3
"""Export WPEC SG52 Exercise 1 GLLS results to NEA-style Excel workbooks.

This script only formats an existing 979-component prior/posterior result. It
performs no GLLS calculation. Edit the configuration constants below if the
repository layout or filenames differ, then run this file directly.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import h5py
import numpy as np
from numpy.typing import NDArray
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# =============================================================================
# Configuration -- edit only this section when paths or filenames change.
# =============================================================================

TEMPLATE_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/SG52_materials_release_2/Output template/Output_example.xlsx")
MU_X_I_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/mu_x_i.hdf5")
SIGMA_X_I_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/Sigma_x_i.hdf5")
DELTA_E_I_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/Delta_E_i.hdf5")
S_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/S.hdf5")
SIGMA_E_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/Sigma_E.hdf5")
POSTERIOR_PATH = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/ex1_results.hdf5")

OUTPUT_DIR = Path("/home/aurja/Programming/OECDNEA_WPEC_SG52/Auroras_solution/Exercise_1/Submission")
RELATIVE_OUTPUT_PATH = OUTPUT_DIR / "Ex1_relative_results.xlsx"
ABSOLUTE_OUTPUT_PATH = OUTPUT_DIR / "Ex1_absolute_results.xlsx"

VERY_LARGE_RELATIVE_THRESHOLD = 1.0e6
# =============================================================================


ZAID = "94239"
N_GROUPS = 51
EXPECTED_STATE_SIZE = 979
SCIENTIFIC_FORMAT = "0.000000000000000E+00"
NAN_TEXT = "NaN"

PRIOR_MTS = [
    "00002",
    "00004",
    "00016",
    "00017",
    "00018",
    "00037",
    "00102",
    "00452",
    "00456",
]
PRIOR_PFNS_IDS = [
    "01018",
    "01118",
    "01218",
    "01318",
    "01418",
    "01518",
    "01618",
    "01718",
    "01818",
    "01918",
]
PRIOR_MUBAR_IDS = [
    "00251",
    "01251",
    "02251",
    "03251",
    "04251",
    "05251",
    "06251",
    "07251",
    "08251",
    "09251",
]

XS_MEAN_HEADERS = [2, 4, 16, 17, 18, 37, 102, 251, 452]
PFNS_MEAN_HEADERS = [1018, 1118, 1218, 1318, 1418, 1518, 1618, 1718, 1818, 1918]
XS_COVARIANCE_MTS = ["00002", "00004", "00016", "00017", "00018", "00037", "00102"]

RELATIVE_SHEET_NAMES = [
    "Comment",
    "XS_posterior_rel_adjust",
    "PFNS_posterior_rel_adjust",
    "XS_cov_posterior_rel_adjust",
    "PFNS_cov_posterior_rel_adjust",
]
ABSOLUTE_SHEET_NAMES = [
    "Comment",
    "State_values_absolute",
    "Covariance_posterior_absolute",
    "Covariance_prior_absolute",
]

FloatArray = NDArray[np.float64]


@dataclass(frozen=True)
class InputData:
    mu_i: FloatArray
    Sigma_i: FloatArray
    Delta_E_i: FloatArray
    S: FloatArray
    Sigma_E: FloatArray
    mu_f: FloatArray
    Sigma_f: FloatArray


@dataclass(frozen=True)
class StateMeta:
    index: int
    label: str
    family: str
    qid: str
    group: int | None
    incident_group: int | None
    outgoing_group: int | None


@dataclass(frozen=True)
class TemplateSelection:
    xs_indices: NDArray[np.int64]
    xs_labels: tuple[str, ...]
    pfns_indices: NDArray[np.int64]
    pfns_labels: tuple[str, ...]

    @property
    def all_indices(self) -> NDArray[np.int64]:
        return np.concatenate((self.xs_indices, self.pfns_indices))


@dataclass(frozen=True)
class RelativeOutputs:
    mean_adjustment: FloatArray
    xs_covariance: FloatArray
    pfns_covariance: FloatArray


@dataclass(frozen=True)
class ExportDiagnostics:
    undefined_mean_count: int
    undefined_covariance_count: int
    very_large_finite_count: int


def read_dataset(path: Path, dataset_path: str) -> FloatArray:
    """Read one HDF5 dataset as float64, with a descriptive missing-data error."""
    if not path.is_file():
        raise FileNotFoundError(f"Required input file does not exist: {path}")

    with h5py.File(path, "r") as h5:
        if dataset_path not in h5:
            raise KeyError(f"{path}: required dataset {dataset_path!r} was not found")
        return np.asarray(h5[dataset_path], dtype=np.float64)


def validate_input_data(data: InputData) -> None:
    expected_shapes = {
        "mu_i": (EXPECTED_STATE_SIZE,),
        "Sigma_i": (EXPECTED_STATE_SIZE, EXPECTED_STATE_SIZE),
        "Delta_E_i": (1,),
        "S": (1, EXPECTED_STATE_SIZE),
        "Sigma_E": (1, 1),
        "mu_f": (EXPECTED_STATE_SIZE,),
        "Sigma_f": (EXPECTED_STATE_SIZE, EXPECTED_STATE_SIZE),
    }

    for name, expected_shape in expected_shapes.items():
        array = getattr(data, name)
        if array.shape != expected_shape:
            raise ValueError(
                f"{name} has shape {array.shape}; expected {expected_shape}. "
                "The supplied files do not match Exercise 1's 979-state layout."
            )
        if not np.all(np.isfinite(array)):
            bad_count = int(array.size - np.count_nonzero(np.isfinite(array)))
            raise ValueError(f"{name} contains {bad_count} non-finite value(s)")


def load_inputs() -> InputData:
    """Load and immediately validate every required numerical input."""
    data = InputData(
        mu_i=read_dataset(MU_X_I_PATH, "/values"),
        Sigma_i=read_dataset(SIGMA_X_I_PATH, "/values"),
        Delta_E_i=read_dataset(DELTA_E_I_PATH, "/values"),
        S=read_dataset(S_PATH, "/values"),
        Sigma_E=read_dataset(SIGMA_E_PATH, "/values"),
        mu_f=read_dataset(POSTERIOR_PATH, "/mean/values"),
        Sigma_f=read_dataset(POSTERIOR_PATH, "/covariance/values"),
    )
    validate_input_data(data)
    return data


def build_state_metadata() -> tuple[list[StateMeta], dict[str, int]]:
    """Reconstruct the canonical Exercise 1 ordering and direct label lookup."""
    metadata: list[StateMeta] = []

    for qid in PRIOR_MTS:
        for group in range(N_GROUPS):
            metadata.append(
                StateMeta(
                    index=len(metadata),
                    label=f"{ZAID}_{qid}_{group}",
                    family="mt",
                    qid=qid,
                    group=group,
                    incident_group=None,
                    outgoing_group=None,
                )
            )

    for incident_group, qid in enumerate(PRIOR_PFNS_IDS):
        for outgoing_group in range(N_GROUPS):
            metadata.append(
                StateMeta(
                    index=len(metadata),
                    label=f"{ZAID}_{qid}_{outgoing_group}",
                    family="pfns",
                    qid=qid,
                    group=outgoing_group,
                    incident_group=incident_group,
                    outgoing_group=outgoing_group,
                )
            )

    for incident_group, qid in enumerate(PRIOR_MUBAR_IDS):
        metadata.append(
            StateMeta(
                index=len(metadata),
                label=f"{ZAID}_{qid}",
                family="mubar",
                qid=qid,
                group=None,
                incident_group=incident_group,
                outgoing_group=None,
            )
        )

    if len(metadata) != EXPECTED_STATE_SIZE:
        raise RuntimeError(
            f"Internal canonical ordering has {len(metadata)} states; "
            f"expected {EXPECTED_STATE_SIZE}"
        )

    index_by_label = {item.label: item.index for item in metadata}
    if len(index_by_label) != EXPECTED_STATE_SIZE:
        raise RuntimeError("Canonical state labels are not unique")
    if any(index_by_label[item.label] != item.index for item in metadata):
        raise RuntimeError("Canonical state label lookup does not preserve ordering")

    return metadata, index_by_label


def build_template_output_indices(index_by_label: dict[str, int]) -> TemplateSelection:
    """Build the exact 418-state XS/mubar and 510-state PFNS selections."""
    xs_labels: list[str] = []
    for qid in XS_COVARIANCE_MTS:
        xs_labels.extend(f"{ZAID}_{qid}_{group}" for group in range(N_GROUPS))
    xs_labels.extend(f"{ZAID}_{qid}" for qid in PRIOR_MUBAR_IDS)
    xs_labels.extend(f"{ZAID}_00452_{group}" for group in range(N_GROUPS))

    pfns_labels = [
        f"{ZAID}_{qid}_{outgoing_group}"
        for qid in PRIOR_PFNS_IDS
        for outgoing_group in range(N_GROUPS)
    ]

    try:
        xs_indices = np.asarray([index_by_label[label] for label in xs_labels], dtype=np.int64)
        pfns_indices = np.asarray(
            [index_by_label[label] for label in pfns_labels], dtype=np.int64
        )
    except KeyError as error:
        raise RuntimeError(f"Template selection references unknown state {error.args[0]!r}") from error

    if xs_indices.shape != (418,):
        raise RuntimeError(f"Official non-PFNS selection has {xs_indices.size} states, expected 418")
    if pfns_indices.shape != (510,):
        raise RuntimeError(f"Official PFNS selection has {pfns_indices.size} states, expected 510")
    if xs_indices.size + pfns_indices.size != 928:
        raise RuntimeError("Official relative selections must contain 418 + 510 = 928 states")

    selected = set(xs_indices.tolist()) | set(pfns_indices.tolist())
    omitted = set(range(EXPECTED_STATE_SIZE)) - selected
    expected_omitted = {
        index_by_label[f"{ZAID}_00456_{group}"] for group in range(N_GROUPS)
    }
    if omitted != expected_omitted:
        raise RuntimeError(
            "The 51 states omitted from the relative template are not exactly MT456"
        )
    if len(selected) != 928:
        raise RuntimeError("Official relative output selections overlap unexpectedly")

    return TemplateSelection(
        xs_indices=xs_indices,
        xs_labels=tuple(xs_labels),
        pfns_indices=pfns_indices,
        pfns_labels=tuple(pfns_labels),
    )


def relative_covariance_block(Sigma_f: FloatArray, mu_f: FloatArray) -> FloatArray:
    """Normalize one selected covariance block without thresholding denominators."""
    denominator = np.multiply.outer(mu_f, mu_f)
    relative = np.full(Sigma_f.shape, np.nan, dtype=np.float64)
    with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
        np.divide(Sigma_f, denominator, out=relative, where=denominator != 0.0)

    if np.any(np.isinf(relative)):
        count = int(np.count_nonzero(np.isinf(relative)))
        raise OverflowError(
            f"Relative covariance calculation produced {count} infinite value(s) "
            "from nonzero denominators; these cannot be represented faithfully in Excel"
        )
    return relative


def calculate_relative_outputs(
    data: InputData, selection: TemplateSelection
) -> RelativeOutputs:
    """Calculate requested relative values exactly, using zero-only guards."""
    mean_adjustment = np.full(data.mu_i.shape, np.nan, dtype=np.float64)
    with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
        np.divide(
            data.mu_f - data.mu_i,
            data.mu_i,
            out=mean_adjustment,
            where=data.mu_i != 0.0,
        )

    if np.any(np.isinf(mean_adjustment)):
        count = int(np.count_nonzero(np.isinf(mean_adjustment)))
        raise OverflowError(
            f"Relative mean calculation produced {count} infinite value(s) from "
            "nonzero denominators; these cannot be represented faithfully in Excel"
        )

    xs_indices = selection.xs_indices
    pfns_indices = selection.pfns_indices
    xs_covariance = relative_covariance_block(
        data.Sigma_f[np.ix_(xs_indices, xs_indices)], data.mu_f[xs_indices]
    )
    pfns_covariance = relative_covariance_block(
        data.Sigma_f[np.ix_(pfns_indices, pfns_indices)], data.mu_f[pfns_indices]
    )

    return RelativeOutputs(
        mean_adjustment=mean_adjustment,
        xs_covariance=xs_covariance,
        pfns_covariance=pfns_covariance,
    )


def assert_allclose_or_raise(
    actual: FloatArray,
    expected: FloatArray,
    description: str,
    *,
    rtol: float = 1.0e-12,
    atol: float = 1.0e-15,
) -> None:
    if not np.allclose(actual, expected, rtol=rtol, atol=atol, equal_nan=True):
        finite_difference = np.abs(actual - expected)
        max_error = float(np.nanmax(finite_difference))
        raise ValueError(f"{description} failed; maximum absolute error is {max_error:.17g}")


def validate_relative_outputs(
    data: InputData,
    selection: TemplateSelection,
    relative: RelativeOutputs,
) -> ExportDiagnostics:
    """Check normalization round trips, ordering, symmetry, and diagnostics."""
    selected_indices = selection.all_indices
    selected_adjustments = relative.mean_adjustment[selected_indices]
    finite_mean = np.isfinite(selected_adjustments)
    reconstructed_mean = data.mu_i[selected_indices][finite_mean] * (
        1.0 + selected_adjustments[finite_mean]
    )
    assert_allclose_or_raise(
        reconstructed_mean,
        data.mu_f[selected_indices][finite_mean],
        "Relative mean round-trip validation",
        rtol=2.0e-12,
        atol=1.0e-12,
    )

    covariance_specs = (
        ("XS", selection.xs_indices, relative.xs_covariance),
        ("PFNS", selection.pfns_indices, relative.pfns_covariance),
    )
    for name, indices, relative_covariance in covariance_specs:
        expected_shape = (indices.size, indices.size)
        if relative_covariance.shape != expected_shape:
            raise ValueError(
                f"{name} relative covariance has shape {relative_covariance.shape}; "
                f"expected {expected_shape}"
            )
        if not np.allclose(
            relative_covariance,
            relative_covariance.T,
            rtol=1.0e-12,
            atol=1.0e-15,
            equal_nan=True,
        ):
            raise ValueError(f"{name} relative covariance is not symmetric after mapping")

        selected_mu = data.mu_f[indices]
        denominator = np.multiply.outer(selected_mu, selected_mu)
        absolute_covariance = data.Sigma_f[np.ix_(indices, indices)]
        finite_covariance = np.isfinite(relative_covariance)
        reconstructed_covariance = (
            relative_covariance[finite_covariance] * denominator[finite_covariance]
        )
        assert_allclose_or_raise(
            reconstructed_covariance,
            absolute_covariance[finite_covariance],
            f"{name} relative covariance round-trip validation",
        )

    covariance_values = (relative.xs_covariance, relative.pfns_covariance)
    undefined_mean_count = int(np.count_nonzero(~finite_mean))
    undefined_covariance_count = sum(
        int(np.count_nonzero(~np.isfinite(values))) for values in covariance_values
    )
    very_large_finite_count = int(
        np.count_nonzero(
            finite_mean
            & (np.abs(selected_adjustments) > VERY_LARGE_RELATIVE_THRESHOLD)
        )
    )
    very_large_finite_count += sum(
        int(
            np.count_nonzero(
                np.isfinite(values)
                & (np.abs(values) > VERY_LARGE_RELATIVE_THRESHOLD)
            )
        )
        for values in covariance_values
    )

    return ExportDiagnostics(
        undefined_mean_count=undefined_mean_count,
        undefined_covariance_count=undefined_covariance_count,
        very_large_finite_count=very_large_finite_count,
    )


def normalized_header(value: Any) -> str:
    """Normalize template header text while preserving exact numeric meaning."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip() if value is not None else ""


def find_unique_cell(ws: Worksheet, expected_value: str) -> Any:
    matches = [
        cell
        for row in ws.iter_rows()
        for cell in row
        if normalized_header(cell.value) == expected_value
    ]
    if len(matches) != 1:
        coordinates = [cell.coordinate for cell in matches]
        raise ValueError(
            f"Sheet {ws.title!r}: expected one {expected_value!r} header, "
            f"found {len(matches)} at {coordinates}"
        )
    return matches[0]


def locate_mean_table(
    ws: Worksheet, anchor_text: str, expected_headers: list[int]
) -> tuple[dict[int, int], dict[int, int]]:
    """Locate mean table columns and group rows from template header values."""
    anchor = find_unique_cell(ws, anchor_text)

    columns: dict[int, int] = {}
    for column in range(anchor.column + 1, ws.max_column + 1):
        value = normalized_header(ws.cell(anchor.row, column).value)
        for expected in expected_headers:
            if value == str(expected):
                if expected in columns:
                    raise ValueError(
                        f"Sheet {ws.title!r}: duplicate column header {expected!r}"
                    )
                columns[expected] = column

    rows: dict[int, int] = {}
    for row in range(anchor.row + 1, ws.max_row + 1):
        value = normalized_header(ws.cell(row, anchor.column).value)
        for group in range(N_GROUPS):
            if value == str(group):
                if group in rows:
                    raise ValueError(
                        f"Sheet {ws.title!r}: duplicate group row {group!r}"
                    )
                rows[group] = row

    missing_headers = [header for header in expected_headers if header not in columns]
    missing_groups = [group for group in range(N_GROUPS) if group not in rows]
    if missing_headers or missing_groups:
        raise ValueError(
            f"Sheet {ws.title!r}: missing headers {missing_headers} or group rows "
            f"{missing_groups}"
        )
    return columns, rows


def locate_covariance_axes(
    ws: Worksheet, expected_labels: tuple[str, ...]
) -> tuple[dict[str, int], dict[str, int]]:
    """Find complete row and column label axes without fixed coordinates."""
    expected = set(expected_labels)
    labels_by_row: dict[int, dict[str, int]] = {}
    labels_by_column: dict[int, dict[str, int]] = {}

    for row in ws.iter_rows():
        for cell in row:
            value = normalized_header(cell.value)
            if value not in expected:
                continue
            row_labels = labels_by_row.setdefault(cell.row, {})
            column_labels = labels_by_column.setdefault(cell.column, {})
            if value in row_labels or value in column_labels:
                raise ValueError(
                    f"Sheet {ws.title!r}: duplicate axis label {value!r} in row "
                    f"{cell.row} or column {cell.column}"
                )
            row_labels[value] = cell.column
            column_labels[value] = cell.row

    row_candidates = [
        mapping for mapping in labels_by_row.values() if set(mapping) == expected
    ]
    column_candidates = [
        mapping for mapping in labels_by_column.values() if set(mapping) == expected
    ]
    if len(row_candidates) != 1 or len(column_candidates) != 1:
        raise ValueError(
            f"Sheet {ws.title!r}: could not identify exactly one complete horizontal "
            f"and vertical covariance axis ({len(row_candidates)} horizontal, "
            f"{len(column_candidates)} vertical)"
        )

    columns = row_candidates[0]
    rows = column_candidates[0]
    if [columns[label] for label in expected_labels] != sorted(columns.values()):
        raise ValueError(f"Sheet {ws.title!r}: horizontal labels are not in required order")
    if [rows[label] for label in expected_labels] != sorted(rows.values()):
        raise ValueError(f"Sheet {ws.title!r}: vertical labels are not in required order")
    return columns, rows


def excel_relative_value(value: float) -> float | str:
    if np.isnan(value):
        return NAN_TEXT
    if not np.isfinite(value):
        raise ValueError(f"Cannot write non-finite relative value {value!r} to Excel")
    return float(value)


def write_relative_cell(cell: Any, value: float) -> None:
    cell.value = excel_relative_value(value)
    if not isinstance(cell.value, str):
        cell.number_format = SCIENTIFIC_FORMAT


def populate_xs_mean_sheet(
    ws: Worksheet,
    adjustment: FloatArray,
    index_by_label: dict[str, int],
) -> None:
    columns, rows = locate_mean_table(ws, "MT", XS_MEAN_HEADERS)

    ordinary_header_qids = {
        2: "00002",
        4: "00004",
        16: "00016",
        17: "00017",
        18: "00018",
        37: "00037",
        102: "00102",
        452: "00452",
    }
    for header, qid in ordinary_header_qids.items():
        for group in range(N_GROUPS):
            label = f"{ZAID}_{qid}_{group}"
            write_relative_cell(
                ws.cell(rows[group], columns[header]),
                adjustment[index_by_label[label]],
            )

    mubar_column = columns[251]
    for incident_group, qid in enumerate(PRIOR_MUBAR_IDS):
        label = f"{ZAID}_{qid}"
        write_relative_cell(
            ws.cell(rows[incident_group], mubar_column),
            adjustment[index_by_label[label]],
        )
    for nonexistent_group in range(len(PRIOR_MUBAR_IDS), N_GROUPS):
        ws.cell(rows[nonexistent_group], mubar_column).value = None


def populate_pfns_mean_sheet(
    ws: Worksheet,
    adjustment: FloatArray,
    index_by_label: dict[str, int],
) -> None:
    columns, rows = locate_mean_table(ws, "Artificial MT", PFNS_MEAN_HEADERS)
    for header, qid in zip(PFNS_MEAN_HEADERS, PRIOR_PFNS_IDS, strict=True):
        for outgoing_group in range(N_GROUPS):
            label = f"{ZAID}_{qid}_{outgoing_group}"
            write_relative_cell(
                ws.cell(rows[outgoing_group], columns[header]),
                adjustment[index_by_label[label]],
            )


def xs_template_axis_labels() -> tuple[str, ...]:
    labels = [
        f"MT{int(qid)}.{group}"
        for qid in XS_COVARIANCE_MTS
        for group in range(N_GROUPS)
    ]
    labels.extend(f"MT251.{incident_group}" for incident_group in range(10))
    labels.extend(f"MT452.{group}" for group in range(N_GROUPS))
    return tuple(labels)


def pfns_template_axis_labels() -> tuple[str, ...]:
    return tuple(
        f"MT{int(qid)}.{outgoing_group}"
        for qid in PRIOR_PFNS_IDS
        for outgoing_group in range(N_GROUPS)
    )


def populate_covariance_sheet(
    ws: Worksheet,
    axis_labels: tuple[str, ...],
    relative_covariance: FloatArray,
) -> None:
    columns, rows = locate_covariance_axes(ws, axis_labels)
    for matrix_row, row_label in enumerate(axis_labels):
        excel_row = rows[row_label]
        for matrix_column, column_label in enumerate(axis_labels):
            write_relative_cell(
                ws.cell(excel_row, columns[column_label]),
                relative_covariance[matrix_row, matrix_column],
            )


def append_relative_comment(ws: Worksheet) -> None:
    start_row = ws.max_row + 2
    lines = [
        "Numerical conventions used for this submission",
        "Relative posterior mean adjustment: a_i = (mu_f,i - mu_i,i) / mu_i,i.",
        "Posterior relative covariance: Sigma_rel,f[i,j] = Sigma_abs,f[i,j] / (mu_f,i * mu_f,j).",
        "Division is performed for every nonzero denominator, regardless of its magnitude.",
        "No epsilon threshold, clipping, or replacement of large finite values was used.",
        "When a mathematical denominator is exactly zero, the relative quantity is undefined and is reported as the literal text NaN.",
        "Blank MT251 cells after incident index 9 represent nonexistent template/state entries; they are not undefined mathematical quantities.",
        "Ex1_absolute_results.xlsx is supplied so the underlying absolute posterior data can be inspected independently of relative normalization.",
        "MT456 and covariance information that cannot be represented by the supplied relative template are retained in the absolute companion workbook.",
    ]
    for offset, line in enumerate(lines):
        cell = ws.cell(start_row + offset, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if offset == 0:
            cell.font = Font(bold=True)


def populate_relative_workbook(
    relative: RelativeOutputs,
    index_by_label: dict[str, int],
) -> None:
    """Load the NEA template, populate mapped cells, and preserve its structure."""
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"Required Excel template does not exist: {TEMPLATE_PATH}")

    workbook = load_workbook(TEMPLATE_PATH)
    missing_sheets = [name for name in RELATIVE_SHEET_NAMES if name not in workbook.sheetnames]
    if missing_sheets:
        workbook.close()
        raise ValueError(f"Template is missing required sheet(s): {missing_sheets}")

    populate_xs_mean_sheet(
        workbook["XS_posterior_rel_adjust"],
        relative.mean_adjustment,
        index_by_label,
    )
    populate_pfns_mean_sheet(
        workbook["PFNS_posterior_rel_adjust"],
        relative.mean_adjustment,
        index_by_label,
    )
    populate_covariance_sheet(
        workbook["XS_cov_posterior_rel_adjust"],
        xs_template_axis_labels(),
        relative.xs_covariance,
    )
    populate_covariance_sheet(
        workbook["PFNS_cov_posterior_rel_adjust"],
        pfns_template_axis_labels(),
        relative.pfns_covariance,
    )
    append_relative_comment(workbook["Comment"])

    workbook.save(RELATIVE_OUTPUT_PATH)
    workbook.close()


def validated_standard_deviation(covariance: FloatArray, name: str) -> FloatArray:
    diagonal = np.diag(covariance)
    if np.any(diagonal < 0.0):
        indices = np.flatnonzero(diagonal < 0.0)
        raise ValueError(
            f"{name} has negative diagonal values at indices {indices[:20].tolist()}"
        )
    return np.sqrt(diagonal)


def write_only_cell(
    ws: Any, value: Any, style_name: str | None = None
) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    if style_name is not None:
        cell.style = style_name
    return cell


def add_absolute_comment_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet("Comment")
    ws.column_dimensions["A"].width = 120
    lines = [
        "Exercise 1 complete absolute companion workbook",
        "This workbook preserves the complete 979-dimensional prior/posterior state and is not the official NEA relative reporting format.",
        "State_values_absolute contains all ordinary-MT values, including MT456, all PFNS values, and all mubar values.",
        "Covariance_posterior_absolute contains the complete posterior absolute covariance matrix with canonical state labels on both axes.",
        "Covariance_prior_absolute contains the complete prior absolute covariance matrix with the same ordering.",
        "The full matrices retain ordinary-MT/PFNS/mubar cross-covariances and every other covariance represented by the HDF5 inputs/results.",
        "The official relative workbook omits exactly the 51 MT456 states because the supplied NEA template has no location for them.",
    ]
    for index, line in enumerate(lines):
        style = "absolute_header" if index == 0 else None
        cell = write_only_cell(ws, line, style)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([cell])


def add_state_values_sheet(
    workbook: Workbook,
    metadata: list[StateMeta],
    data: InputData,
) -> None:
    ws = workbook.create_sheet("State_values_absolute")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{EXPECTED_STATE_SIZE + 1}"

    widths = [10, 24, 12, 10, 10, 16, 16, 21, 21, 21, 21, 21]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    headers = [
        "index",
        "label",
        "family",
        "qid",
        "group",
        "incident_group",
        "outgoing_group",
        "mu_i",
        "mu_f",
        "delta_mu",
        "prior_std",
        "posterior_std",
    ]
    ws.append([write_only_cell(ws, header, "absolute_header") for header in headers])

    prior_std = validated_standard_deviation(data.Sigma_i, "Sigma_i")
    posterior_std = validated_standard_deviation(data.Sigma_f, "Sigma_f")
    delta_mu = data.mu_f - data.mu_i

    for item in metadata:
        numerical_values = (
            data.mu_i[item.index],
            data.mu_f[item.index],
            delta_mu[item.index],
            prior_std[item.index],
            posterior_std[item.index],
        )
        row = [
            item.index,
            item.label,
            item.family,
            item.qid,
            item.group,
            item.incident_group,
            item.outgoing_group,
        ]
        row.extend(
            write_only_cell(ws, float(value), "absolute_scientific")
            for value in numerical_values
        )
        ws.append(row)


def add_absolute_covariance_sheet(
    workbook: Workbook,
    sheet_name: str,
    labels: list[str],
    covariance: FloatArray,
) -> None:
    ws = workbook.create_sheet(sheet_name)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 24

    header = [write_only_cell(ws, "label", "absolute_header")]
    header.extend(
        write_only_cell(ws, label, "absolute_header") for label in labels
    )
    ws.append(header)

    for label, covariance_row in zip(labels, covariance, strict=True):
        row = [label]
        row.extend(
            write_only_cell(ws, float(value), "absolute_scientific")
            for value in covariance_row
        )
        ws.append(row)


def populate_absolute_workbook(metadata: list[StateMeta], data: InputData) -> None:
    """Create the complete absolute workbook in constant-memory write-only mode."""
    workbook = Workbook(write_only=True)
    workbook.add_named_style(
        NamedStyle(name="absolute_header", font=Font(bold=True))
    )
    workbook.add_named_style(
        NamedStyle(name="absolute_scientific", number_format=SCIENTIFIC_FORMAT)
    )

    add_absolute_comment_sheet(workbook)
    add_state_values_sheet(workbook, metadata, data)
    labels = [item.label for item in metadata]
    add_absolute_covariance_sheet(
        workbook,
        "Covariance_posterior_absolute",
        labels,
        data.Sigma_f,
    )
    add_absolute_covariance_sheet(
        workbook,
        "Covariance_prior_absolute",
        labels,
        data.Sigma_i,
    )

    workbook.save(ABSOLUTE_OUTPUT_PATH)
    workbook.close()


def validate_saved_workbook(path: Path, expected_sheets: list[str]) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Expected output workbook was not created: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        missing = [name for name in expected_sheets if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"{path}: saved workbook is missing sheet(s) {missing}")
    finally:
        workbook.close()


def validate_outputs() -> None:
    """Reopen both generated files once and verify their required sheet names."""
    validate_saved_workbook(RELATIVE_OUTPUT_PATH, RELATIVE_SHEET_NAMES)
    validate_saved_workbook(ABSOLUTE_OUTPUT_PATH, ABSOLUTE_SHEET_NAMES)


def main() -> None:
    data = load_inputs()
    metadata, index_by_label = build_state_metadata()
    selection = build_template_output_indices(index_by_label)
    relative = calculate_relative_outputs(data, selection)
    diagnostics = validate_relative_outputs(data, selection, relative)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    populate_relative_workbook(relative, index_by_label)
    populate_absolute_workbook(metadata, data)
    validate_outputs()

    print(f"Loaded {EXPECTED_STATE_SIZE}-state Exercise-1 result.")
    print("Relative template state coverage: 928 / 979.")
    print("Omitted from NEA relative template: MT456 (51 values).")
    print(f"Wrote: {RELATIVE_OUTPUT_PATH}")
    print(f"Wrote: {ABSOLUTE_OUTPUT_PATH}")
    print(
        "Undefined relative mean cells written as NaN: "
        f"{diagnostics.undefined_mean_count}"
    )
    print(
        "Undefined relative covariance cells written as NaN: "
        f"{diagnostics.undefined_covariance_count}"
    )
    print(
        "Very large finite relative values "
        f"(|value| > {VERY_LARGE_RELATIVE_THRESHOLD:.1e}): "
        f"{diagnostics.very_large_finite_count}"
    )
    print("Export validation passed.")


if __name__ == "__main__":
    main()
