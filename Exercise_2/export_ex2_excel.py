#!/usr/bin/env python3
"""Export WPEC SG52 Exercise 2 GLLS results to NEA-style Excel workbooks.

This script only formats an existing 979-component prior/posterior result. It
performs no GLLS calculation. Edit the configuration constants below if the
repository layout or filenames differ, then run this file directly.
"""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from numbers import Integral, Real
from pathlib import Path
from typing import cast

import h5py
import numpy as np
from numpy.typing import NDArray
from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell, WriteOnlyCell
from openpyxl.styles import Alignment, Font, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# =============================================================================
# Configuration -- edit only this section when paths or filenames change.
# =============================================================================

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
EXERCISE_DIRECTORY = Path(__file__).resolve().parent
TEMPLATE_PATH = (
    REPOSITORY_ROOT.parent
    / "SG52_materials"
    / "SG52_materials_release_3"
    / "Output template"
    / "Output_example.xlsx"
)
PRIOR_PATH = EXERCISE_DIRECTORY / "ex2_prior.hdf5"
OBSERVATION_PATH = EXERCISE_DIRECTORY / "ex2_observations.hdf5"
POSTERIOR_PATH = EXERCISE_DIRECTORY / "ex2_results.hdf5"

OUTPUT_DIR = EXERCISE_DIRECTORY / "Submission"
RELATIVE_OUTPUT_PATH = OUTPUT_DIR / "Ex2_relative_results.xlsx"
ABSOLUTE_OUTPUT_PATH = OUTPUT_DIR / "Ex2_absolute_results.xlsx"

VERY_LARGE_RELATIVE_THRESHOLD = 1.0e6
# =============================================================================


ZAID = "94239"
N_GROUPS = 51
EXPECTED_STATE_SIZE = 979
SCIENTIFIC_FORMAT = "0.0000000000E+00"
NAN_TEXT = "NaN"

PRIOR_MTS = [
    "00002",
    "00004",
    "00016",
    "00017",
    "00018",
    "00037",
    "00102",
    "00452",
    "00456",
]
PRIOR_PFNS_IDS = [
    "01018",
    "01118",
    "01218",
    "01318",
    "01418",
    "01518",
    "01618",
    "01718",
    "01818",
    "01918",
]
PRIOR_MUBAR_IDS = [
    "00251",
    "01251",
    "02251",
    "03251",
    "04251",
    "05251",
    "06251",
    "07251",
    "08251",
    "09251",
]

MEAN_HEADERS = [
    2,
    4,
    16,
    17,
    18,
    37,
    102,
    251,
    452,
    1018,
    1118,
    1218,
    1318,
    1418,
    1518,
    1618,
    1718,
    1818,
    1918,
]
MUBAR_HEADER = 251
MEAN_TABLE_ANCHOR = "GROUP\\MT"
TEMPLATE_STATE_COUNT = 928

RELATIVE_SHEET_NAMES = [
    "Comment",
    "Questions",
    "Groups",
    "prior_abs",
    "prior_cov_rel_to_prior",
    "posterior_rel_adjust",
    "posterior_cov_rel_to_prior",
]
ABSOLUTE_SHEET_NAMES = [
    "Comment",
    "State_values_absolute",
    "Covariance_posterior_absolute",
    "Covariance_prior_absolute",
]

FloatArray = NDArray[np.float64]
BoolArray = NDArray[np.bool_]
ExcelCellValue = str | int | float | None
ExcelRowValue = Cell | ExcelCellValue


def float_at(values: FloatArray, index: int) -> float:
    return float(cast(np.float64, values[index]))


def float_at_2d(values: FloatArray, row: int, column: int) -> float:
    return float(cast(np.float64, values[row, column]))


@dataclass(frozen=True)
class InputData:
    mu_i: FloatArray
    Sigma_i: FloatArray
    Delta_E_i: FloatArray
    S: FloatArray
    Sigma_E: FloatArray
    mu_f: FloatArray
    Sigma_f: FloatArray


@dataclass(frozen=True)
class StateMeta:
    index: int
    label: str
    family: str
    qid: str
    group: int | None
    incident_group: int | None
    outgoing_group: int | None


@dataclass(frozen=True)
class TemplateSelection:
    indices: NDArray[np.int64]
    axis_labels: tuple[str, ...]


@dataclass(frozen=True)
class RelativeOutputs:
    mean_adjustment: FloatArray
    prior_covariance: FloatArray
    posterior_covariance: FloatArray


@dataclass(frozen=True)
class ExportDiagnostics:
    undefined_mean_count: int
    undefined_prior_covariance_count: int
    undefined_posterior_covariance_count: int
    very_large_finite_count: int


def read_dataset(path: Path, dataset_path: str) -> FloatArray:
    """Read one HDF5 dataset as float64, with a descriptive missing-data error."""
    if not path.is_file():
        raise FileNotFoundError(f"Required input file does not exist: {path}")

    with h5py.File(str(path), "r") as h5:
        if dataset_path not in h5:
            raise KeyError(f"{path}: required dataset {dataset_path!r} was not found")
        dataset = cast(h5py.Dataset, h5[dataset_path])
        values: FloatArray = np.empty(dataset.shape, dtype=np.float64)
        dataset.read_direct(values)
        return values


def validate_input_data(data: InputData) -> None:
    arrays: dict[str, FloatArray] = {
        "mu_i": data.mu_i,
        "Sigma_i": data.Sigma_i,
        "Delta_E_i": data.Delta_E_i,
        "S": data.S,
        "Sigma_E": data.Sigma_E,
        "mu_f": data.mu_f,
        "Sigma_f": data.Sigma_f,
    }
    expected_shapes: dict[str, tuple[int, ...]] = {
        "mu_i": (EXPECTED_STATE_SIZE,),
        "Sigma_i": (EXPECTED_STATE_SIZE, EXPECTED_STATE_SIZE),
        "Delta_E_i": (3,),
        "S": (3, EXPECTED_STATE_SIZE),
        "Sigma_E": (3, 3),
        "mu_f": (EXPECTED_STATE_SIZE,),
        "Sigma_f": (EXPECTED_STATE_SIZE, EXPECTED_STATE_SIZE),
    }

    for name, array in arrays.items():
        expected_shape = expected_shapes[name]
        if array.shape != expected_shape:
            raise ValueError(
                f"{name} has shape {array.shape}; expected {expected_shape}. "
                + "The supplied files do not match Exercise 2's three-response, 979-state layout."
            )
        if not np.all(np.isfinite(array)):
            bad_count = int(array.size - np.count_nonzero(np.isfinite(array)))
            raise ValueError(f"{name} contains {bad_count} non-finite value(s)")


def load_inputs() -> InputData:
    """Load and immediately validate every required numerical input."""
    data = InputData(
        mu_i=read_dataset(PRIOR_PATH, "/mean/values"),
        Sigma_i=read_dataset(PRIOR_PATH, "/cov/values"),
        Delta_E_i=read_dataset(OBSERVATION_PATH, "/delta/values"),
        S=read_dataset(OBSERVATION_PATH, "/sens/values"),
        Sigma_E=read_dataset(OBSERVATION_PATH, "/cov/values"),
        mu_f=read_dataset(POSTERIOR_PATH, "/mean/values"),
        Sigma_f=read_dataset(POSTERIOR_PATH, "/covariance/values"),
    )
    validate_input_data(data)
    return data


def build_state_metadata() -> tuple[list[StateMeta], dict[str, int]]:
    """Reconstruct the canonical Pu-239 ordering and direct label lookup."""
    metadata: list[StateMeta] = []

    for qid in PRIOR_MTS:
        for group in range(N_GROUPS):
            metadata.append(
                StateMeta(
                    index=len(metadata),
                    label=f"{ZAID}_{qid}_{group}",
                    family="mt",
                    qid=qid,
                    group=group,
                    incident_group=None,
                    outgoing_group=None,
                )
            )

    for incident_group, qid in enumerate(PRIOR_PFNS_IDS):
        for outgoing_group in range(N_GROUPS):
            metadata.append(
                StateMeta(
                    index=len(metadata),
                    label=f"{ZAID}_{qid}_{outgoing_group}",
                    family="pfns",
                    qid=qid,
                    group=outgoing_group,
                    incident_group=incident_group,
                    outgoing_group=outgoing_group,
                )
            )

    for incident_group, qid in enumerate(PRIOR_MUBAR_IDS):
        metadata.append(
            StateMeta(
                index=len(metadata),
                label=f"{ZAID}_{qid}",
                family="mubar",
                qid=qid,
                group=None,
                incident_group=incident_group,
                outgoing_group=None,
            )
        )

    if len(metadata) != EXPECTED_STATE_SIZE:
        raise RuntimeError(
            f"Internal canonical ordering has {len(metadata)} states; "
            + f"expected {EXPECTED_STATE_SIZE}"
        )

    index_by_label = {item.label: item.index for item in metadata}
    if len(index_by_label) != EXPECTED_STATE_SIZE:
        raise RuntimeError("Canonical state labels are not unique")
    if any(index_by_label[item.label] != item.index for item in metadata):
        raise RuntimeError("Canonical state label lookup does not preserve ordering")

    return metadata, index_by_label


def template_cells() -> list[tuple[int, int, str]]:
    """Return the 928 NEA template cells as (MT header, group, state label)."""
    cells: list[tuple[int, int, str]] = []
    for header in MEAN_HEADERS:
        if header == MUBAR_HEADER:
            cells.extend(
                (header, group, f"{ZAID}_{qid}")
                for group, qid in enumerate(PRIOR_MUBAR_IDS)
            )
        else:
            cells.extend(
                (header, group, f"{ZAID}_{header:05d}_{group}")
                for group in range(N_GROUPS)
            )
    return cells


def build_template_selection(index_by_label: dict[str, int]) -> TemplateSelection:
    """Map the 928 NEA template cells to canonical state indices, in template order."""
    cells = template_cells()
    try:
        indices: NDArray[np.int64] = np.asarray(
            [index_by_label[label] for _, _, label in cells], dtype=np.int64
        )
    except KeyError as error:
        raise RuntimeError(
            f"Template selection references unknown state {error.args[0]!r}"
        ) from error

    if indices.shape != (TEMPLATE_STATE_COUNT,):
        raise RuntimeError(
            f"Template selection has {indices.size} states, "
            + f"expected {TEMPLATE_STATE_COUNT}"
        )
    unique_indices = set(indices.tolist())
    if len(unique_indices) != TEMPLATE_STATE_COUNT:
        raise RuntimeError("Template selection contains duplicate states")

    omitted = set(range(EXPECTED_STATE_SIZE)) - unique_indices
    expected_omitted = {
        index_by_label[f"{ZAID}_00456_{group}"] for group in range(N_GROUPS)
    }
    if omitted != expected_omitted:
        raise RuntimeError(
            "The 51 states omitted from the relative template are not exactly MT456"
        )

    return TemplateSelection(
        indices=indices,
        axis_labels=tuple(f"MT{header}.{group}" for header, group, _ in cells),
    )


def relative_covariance_block(
    covariance: FloatArray, prior_mean: FloatArray
) -> FloatArray:
    """Normalize one selected covariance block by the prior-mean outer product."""
    denominator: FloatArray = cast(
        FloatArray, np.multiply.outer(prior_mean, prior_mean)
    )
    relative: FloatArray = np.full(covariance.shape, np.nan, dtype=np.float64)
    denominator_nonzero: BoolArray = cast(BoolArray, denominator != 0.0)
    with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
        _ = cast(
            FloatArray,
            np.divide(
                covariance,
                denominator,
                out=relative,
                where=denominator_nonzero,
            ),
        )

    if np.any(np.isinf(relative)):
        count = int(np.count_nonzero(np.isinf(relative)))
        raise OverflowError(
            f"Relative covariance calculation produced {count} infinite value(s) "
            + "from nonzero denominators; these cannot be represented faithfully in Excel"
        )
    return relative


def calculate_relative_outputs(
    data: InputData, selection: TemplateSelection
) -> RelativeOutputs:
    """Normalize all relative quantities on the prior mean mu_i, not the posterior."""
    mean_adjustment: FloatArray = np.full(data.mu_i.shape, np.nan, dtype=np.float64)
    nonzero_prior_mean: BoolArray = cast(BoolArray, data.mu_i != 0.0)
    with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
        _ = cast(
            FloatArray,
            np.divide(
                data.mu_f - data.mu_i,
                data.mu_i,
                out=mean_adjustment,
                where=nonzero_prior_mean,
            ),
        )

    if np.any(np.isinf(mean_adjustment)):
        count = int(np.count_nonzero(np.isinf(mean_adjustment)))
        raise OverflowError(
            f"Relative mean calculation produced {count} infinite value(s) from "
            + "nonzero denominators; these cannot be represented faithfully in Excel"
        )

    selected = selection.indices
    prior_covariance = relative_covariance_block(
        data.Sigma_i[np.ix_(selected, selected)], data.mu_i[selected]
    )
    posterior_covariance = relative_covariance_block(
        data.Sigma_f[np.ix_(selected, selected)], data.mu_i[selected]
    )

    return RelativeOutputs(
        mean_adjustment=mean_adjustment,
        prior_covariance=prior_covariance,
        posterior_covariance=posterior_covariance,
    )


def assert_allclose_or_raise(
    actual: FloatArray,
    expected: FloatArray,
    description: str,
    *,
    rtol: float = 1.0e-12,
    atol: float = 1.0e-15,
) -> None:
    if not np.allclose(actual, expected, rtol=rtol, atol=atol, equal_nan=True):
        finite_difference: FloatArray = cast(FloatArray, np.abs(actual - expected))
        max_error = float(np.nanmax(finite_difference))
        raise ValueError(
            f"{description} failed; maximum absolute error is {max_error:.17g}"
        )


def validate_relative_outputs(
    data: InputData,
    selection: TemplateSelection,
    relative: RelativeOutputs,
) -> ExportDiagnostics:
    """Check mean and covariance round trips, symmetry, and undefined counts."""
    selected = selection.indices
    prior_mean = data.mu_i[selected]
    selected_adjustments = relative.mean_adjustment[selected]
    finite_mean = np.isfinite(selected_adjustments)
    reconstructed_mean = prior_mean[finite_mean] * (
        1.0 + selected_adjustments[finite_mean]
    )
    assert_allclose_or_raise(
        reconstructed_mean,
        data.mu_f[selected][finite_mean],
        "Relative mean round-trip validation",
        rtol=2.0e-12,
        atol=1.0e-12,
    )

    denominator: FloatArray = cast(
        FloatArray, np.multiply.outer(prior_mean, prior_mean)
    )
    covariance_specs = (
        ("prior", data.Sigma_i, relative.prior_covariance),
        ("posterior", data.Sigma_f, relative.posterior_covariance),
    )
    for name, absolute_matrix, relative_covariance in covariance_specs:
        expected_shape = (selected.size, selected.size)
        if relative_covariance.shape != expected_shape:
            raise ValueError(
                f"{name} relative covariance has shape {relative_covariance.shape}; "
                + f"expected {expected_shape}"
            )
        if not np.allclose(
            relative_covariance,
            relative_covariance.T,
            rtol=1.0e-12,
            atol=1.0e-15,
            equal_nan=True,
        ):
            raise ValueError(
                f"{name} relative covariance is not symmetric after mapping"
            )

        absolute_covariance = absolute_matrix[np.ix_(selected, selected)]
        finite_covariance = np.isfinite(relative_covariance)
        reconstructed_covariance = (
            relative_covariance[finite_covariance] * denominator[finite_covariance]
        )
        assert_allclose_or_raise(
            reconstructed_covariance,
            absolute_covariance[finite_covariance],
            f"{name} relative covariance round-trip validation",
        )

    undefined_mean_count = int(np.count_nonzero(~finite_mean))
    undefined_prior_covariance_count = int(
        np.count_nonzero(~np.isfinite(relative.prior_covariance))
    )
    undefined_posterior_covariance_count = int(
        np.count_nonzero(~np.isfinite(relative.posterior_covariance))
    )
    very_large_finite_count = int(
        np.count_nonzero(
            finite_mean & (np.abs(selected_adjustments) > VERY_LARGE_RELATIVE_THRESHOLD)
        )
    )
    very_large_finite_count += sum(
        int(
            np.count_nonzero(
                np.isfinite(values) & (np.abs(values) > VERY_LARGE_RELATIVE_THRESHOLD)
            )
        )
        for values in (relative.prior_covariance, relative.posterior_covariance)
    )

    return ExportDiagnostics(
        undefined_mean_count=undefined_mean_count,
        undefined_prior_covariance_count=undefined_prior_covariance_count,
        undefined_posterior_covariance_count=undefined_posterior_covariance_count,
        very_large_finite_count=very_large_finite_count,
    )


def normalized_header(value: object) -> str:
    """Normalize template header text while preserving exact numeric meaning."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, Integral):
        return str(int(value))
    if isinstance(value, Real):
        numeric_value = float(value)
        if numeric_value.is_integer():
            return str(int(numeric_value))
    return str(value).strip() if value is not None else ""


def find_unique_cell(ws: Worksheet, expected_value: str) -> Cell:
    matches: list[Cell] = [
        cell
        for row in cast(Iterable[tuple[Cell, ...]], ws.iter_rows())
        for cell in row
        if normalized_header(cast(object, cell.value)) == expected_value
    ]
    if len(matches) != 1:
        coordinates = [cell.coordinate for cell in matches]
        raise ValueError(
            f"Sheet {ws.title!r}: expected one {expected_value!r} header, "
            + f"found {len(matches)} at {coordinates}"
        )
    return matches[0]


def locate_mean_table(
    ws: Worksheet, anchor_text: str, expected_headers: list[int]
) -> tuple[dict[int, int], dict[int, int]]:
    """Locate mean table columns and group rows from template header values."""
    anchor = find_unique_cell(ws, anchor_text)
    max_column = ws.max_column
    max_row = ws.max_row
    columns: dict[int, int] = {}
    for column in range(anchor.column + 1, max_column + 1):
        value = normalized_header(cast(object, ws.cell(anchor.row, column).value))
        for expected in expected_headers:
            if value == str(expected):
                if expected in columns:
                    raise ValueError(
                        f"Sheet {ws.title!r}: duplicate column header {expected!r}"
                    )
                columns[expected] = column

    rows: dict[int, int] = {}
    for row in range(anchor.row + 1, max_row + 1):
        value = normalized_header(cast(object, ws.cell(row, anchor.column).value))
        for group in range(N_GROUPS):
            if value == str(group):
                if group in rows:
                    raise ValueError(
                        f"Sheet {ws.title!r}: duplicate group row {group!r}"
                    )
                rows[group] = row

    missing_headers = [header for header in expected_headers if header not in columns]
    missing_groups = [group for group in range(N_GROUPS) if group not in rows]
    if missing_headers or missing_groups:
        raise ValueError(
            f"Sheet {ws.title!r}: missing headers {missing_headers} or group rows "
            + f"{missing_groups}"
        )
    return columns, rows


def locate_covariance_axes(
    ws: Worksheet, expected_labels: tuple[str, ...]
) -> tuple[dict[str, int], dict[str, int]]:
    """Find complete row and column label axes without fixed coordinates."""
    expected = set(expected_labels)
    labels_by_row: dict[int, dict[str, int]] = {}
    labels_by_column: dict[int, dict[str, int]] = {}

    for row in cast(Iterable[tuple[Cell, ...]], ws.iter_rows()):
        for cell in row:
            value = normalized_header(cast(object, cell.value))
            if value not in expected:
                continue
            row_labels = labels_by_row.setdefault(cell.row, {})
            column_labels = labels_by_column.setdefault(cell.column, {})
            if value in row_labels or value in column_labels:
                raise ValueError(
                    f"Sheet {ws.title!r}: duplicate axis label {value!r} in row "
                    + f"{cell.row} or column {cell.column}"
                )
            row_labels[value] = cell.column
            column_labels[value] = cell.row

    row_candidates = [
        mapping for mapping in labels_by_row.values() if set(mapping) == expected
    ]
    column_candidates = [
        mapping for mapping in labels_by_column.values() if set(mapping) == expected
    ]
    if len(row_candidates) != 1 or len(column_candidates) != 1:
        raise ValueError(
            f"Sheet {ws.title!r}: could not identify exactly one complete horizontal "
            + f"and vertical covariance axis ({len(row_candidates)} horizontal, "
            + f"{len(column_candidates)} vertical)"
        )

    columns = row_candidates[0]
    rows = column_candidates[0]
    if [columns[label] for label in expected_labels] != sorted(columns.values()):
        raise ValueError(
            f"Sheet {ws.title!r}: horizontal labels are not in required order"
        )
    if [rows[label] for label in expected_labels] != sorted(rows.values()):
        raise ValueError(
            f"Sheet {ws.title!r}: vertical labels are not in required order"
        )
    return columns, rows


def excel_numeric_value(value: float) -> float | str:
    if np.isnan(value):
        return NAN_TEXT
    if not np.isfinite(value):
        raise ValueError(f"Cannot write non-finite value {value!r} to Excel")
    return float(value)


def write_numeric_cell(cell: Cell, value: float) -> None:
    cell.value = excel_numeric_value(value)
    if not isinstance(cell.value, str):
        cell.number_format = SCIENTIFIC_FORMAT


def populate_mean_sheet(
    ws: Worksheet,
    values: FloatArray,
    index_by_label: dict[str, int],
) -> None:
    """Fill one MT-by-group mean table from a 979-state vector."""
    columns, rows = locate_mean_table(ws, MEAN_TABLE_ANCHOR, MEAN_HEADERS)
    for header, group, label in template_cells():
        write_numeric_cell(
            cast(Cell, ws.cell(rows[group], columns[header])),
            float_at(values, index_by_label[label]),
        )


def populate_covariance_sheet(
    ws: Worksheet,
    axis_labels: tuple[str, ...],
    covariance: FloatArray,
) -> None:
    columns, rows = locate_covariance_axes(ws, axis_labels)
    for matrix_row, row_label in enumerate(axis_labels):
        excel_row = rows[row_label]
        for matrix_column, column_label in enumerate(axis_labels):
            write_numeric_cell(
                cast(Cell, ws.cell(excel_row, columns[column_label])),
                float_at_2d(covariance, matrix_row, matrix_column),
            )


RELATIVE_COMMENT_LINES = (
    "Numerical conventions used for this submission",
    "Relative mean adjustment: a_i = (mu_post,i - mu_prior,i) / mu_prior,i.",
    "Prior and posterior relative covariances: Sigma[i,j] / (mu_prior,i * mu_prior,j), normalized on the prior mean.",
    "Division is performed for every nonzero denominator, regardless of its magnitude; no thresholds or clipping are applied.",
    "When a denominator is exactly zero, the relative quantity is undefined and is reported as the literal text NaN.",
    "Blank MT251 cells after incident index 9 are not states; they are not undefined mathematical quantities.",
    "prior_abs holds the absolute prior mean mu_prior,i as ingested.",
    "Numbers are shown in scientific format with 11 significant digits.",
    "MT456 (51 states) is not in this template; its full absolute covariances are kept in Ex2_absolute_results.xlsx.",
)


def append_relative_comment(ws: Worksheet) -> None:
    start_row = ws.max_row + 2
    for offset, line in enumerate(RELATIVE_COMMENT_LINES):
        cell = ws.cell(start_row + offset, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if offset == 0:
            cell.font = Font(bold=True)


QUESTION_ANSWERS = (
    (
        "Are you considering the input covariances as relative covariances ?",
        "Yes. The released covariances are relative. Each entry is multiplied by the two prior mean values, Sigma_abs[i,j] = mu_prior,i * mu_prior,j * R[i,j], to form the absolute prior covariance used in GLLS.",
    ),
    (
        "Is your code using GLLS equations in relative or absolute form ?",
        "Absolute form: mu_post = mu_prior + Sigma_prior S^T V^-1 (E - C), Sigma_post = Sigma_prior - Sigma_prior S^T V^-1 S Sigma_prior, with V = S Sigma_prior S^T + Sigma_E.",
    ),
    (
        "If your code uses the relative form of GLLS, are all the terms (bias, experimental covariance, model/calculation covariance) normalized on C ?",
        "Not applicable: GLLS is solved in absolute form. All terms are absolute: discrepancy E - C, experimental/calculation covariance Sigma_E = diag(expt_uncert^2 + sim_uncert^2), and the absolute prior covariance. Relative sensitivities are converted to absolute ones with the calculated value C: S[k,j] = C_k * s_rel[k,j] / mu_prior,j.",
    ),
    (
        "Are your posterior covariances normalized on the prior mean to match the exercise specifications ?",
        "Yes. Both relative covariance sheets are divided by the prior means: Sigma[i,j] / (mu_prior,i * mu_prior,j).",
    ),
    (
        "Are you taking into account the sensitivity uncertainties ? If yes, elaborate briefly on how you do it.",
        "No. The sensitivity uncertainties (sens_unc_rel.csv) are not propagated; the sensitivities are treated as exact.",
    ),
)


def question_row(ws: Worksheet, question: str) -> int:
    """Return the single row whose column-A text, whitespace-stripped, equals question."""
    rows = [
        row[0].row
        for row in cast(Iterable[tuple[Cell, ...]], ws.iter_rows(min_col=1, max_col=1))
        if normalized_header(cast(object, row[0].value)) == question
    ]
    if len(rows) != 1:
        raise ValueError(
            f"Sheet {ws.title!r}: expected one row for question {question!r}, "
            + f"found {len(rows)}"
        )
    return rows[0]


def populate_questions_sheet(ws: Worksheet) -> None:
    for question, answer in QUESTION_ANSWERS:
        answer_cell = cast(Cell, ws.cell(question_row(ws, question), 2))
        answer_cell.value = answer
        answer_cell.alignment = Alignment(wrap_text=True, vertical="top")


def populate_relative_workbook(
    data: InputData,
    relative: RelativeOutputs,
    selection: TemplateSelection,
    index_by_label: dict[str, int],
) -> None:
    """Load the NEA release-3 template, populate mapped cells, and preserve its structure."""
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(
            f"Required Excel template does not exist: {TEMPLATE_PATH}"
        )

    workbook = load_workbook(str(TEMPLATE_PATH))
    missing_sheets = [
        name for name in RELATIVE_SHEET_NAMES if name not in workbook.sheetnames
    ]
    if missing_sheets:
        workbook.close()
        raise ValueError(f"Template is missing required sheet(s): {missing_sheets}")

    populate_questions_sheet(cast(Worksheet, workbook["Questions"]))
    populate_mean_sheet(
        cast(Worksheet, workbook["prior_abs"]), data.mu_i, index_by_label
    )
    populate_mean_sheet(
        cast(Worksheet, workbook["posterior_rel_adjust"]),
        relative.mean_adjustment,
        index_by_label,
    )
    populate_covariance_sheet(
        cast(Worksheet, workbook["prior_cov_rel_to_prior"]),
        selection.axis_labels,
        relative.prior_covariance,
    )
    populate_covariance_sheet(
        cast(Worksheet, workbook["posterior_cov_rel_to_prior"]),
        selection.axis_labels,
        relative.posterior_covariance,
    )
    append_relative_comment(cast(Worksheet, workbook["Comment"]))

    workbook.save(RELATIVE_OUTPUT_PATH)
    workbook.close()


def validated_standard_deviation(covariance: FloatArray, name: str) -> FloatArray:
    diagonal = np.diag(covariance)
    if np.any(diagonal < 0.0):
        indices = np.flatnonzero(diagonal < 0.0)
        raise ValueError(
            f"{name} has negative diagonal values at indices {indices[:20].tolist()}"
        )
    return np.sqrt(diagonal)


def write_only_cell(
    ws: Worksheet, value: ExcelCellValue, style_name: str | None = None
) -> Cell:
    cell = WriteOnlyCell(ws)
    cell.value = value
    if style_name is not None:
        cell.style = style_name
    return cell


def add_absolute_comment_sheet(workbook: Workbook) -> None:
    ws = cast(Worksheet, workbook.create_sheet("Comment"))
    ws.column_dimensions["A"].width = 120
    lines = [
        "Exercise 2 complete absolute companion workbook",
        "This workbook preserves the complete 979-dimensional prior/posterior state and is not the official NEA relative reporting format.",
        "State_values_absolute contains all ordinary-MT values, including MT456, all PFNS values, and all mubar values.",
        "Covariance_posterior_absolute contains the complete posterior absolute covariance matrix with canonical state labels on both axes.",
        "Covariance_prior_absolute contains the complete prior absolute covariance matrix with the same ordering.",
        "The full matrices retain ordinary-MT/PFNS/mubar cross-covariances and every other covariance represented by the HDF5 inputs/results.",
        "The official relative workbook omits exactly the 51 MT456 states because the supplied NEA template has no location for them.",
    ]
    for index, line in enumerate(lines):
        style = "absolute_header" if index == 0 else None
        cell = write_only_cell(ws, line, style)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([cell])


def add_state_values_sheet(
    workbook: Workbook,
    metadata: list[StateMeta],
    data: InputData,
) -> None:
    ws = cast(Worksheet, workbook.create_sheet("State_values_absolute"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{EXPECTED_STATE_SIZE + 1}"

    widths = [10, 24, 12, 10, 10, 16, 16, 21, 21, 21, 21, 21]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    headers = [
        "index",
        "label",
        "family",
        "qid",
        "group",
        "incident_group",
        "outgoing_group",
        "mu_i",
        "mu_f",
        "delta_mu",
        "prior_std",
        "posterior_std",
    ]
    ws.append([write_only_cell(ws, header, "absolute_header") for header in headers])

    prior_std = validated_standard_deviation(data.Sigma_i, "Sigma_i")
    posterior_std = validated_standard_deviation(data.Sigma_f, "Sigma_f")
    delta_mu = data.mu_f - data.mu_i

    for item in metadata:
        numerical_values = (
            float_at(data.mu_i, item.index),
            float_at(data.mu_f, item.index),
            float_at(delta_mu, item.index),
            float_at(prior_std, item.index),
            float_at(posterior_std, item.index),
        )
        row: list[ExcelRowValue] = [
            item.index,
            item.label,
            item.family,
            item.qid,
            item.group,
            item.incident_group,
            item.outgoing_group,
        ]
        row.extend(
            write_only_cell(ws, float(value), "absolute_scientific")
            for value in numerical_values
        )
        ws.append(row)


def add_absolute_covariance_sheet(
    workbook: Workbook,
    sheet_name: str,
    labels: list[str],
    covariance: FloatArray,
) -> None:
    ws = cast(Worksheet, workbook.create_sheet(sheet_name))
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 24

    header = [write_only_cell(ws, "label", "absolute_header")]
    header.extend(write_only_cell(ws, label, "absolute_header") for label in labels)
    ws.append(header)

    for row_index, label in enumerate(labels):
        row: list[ExcelRowValue] = [label]
        row.extend(
            write_only_cell(
                ws,
                float_at_2d(covariance, row_index, column),
                "absolute_scientific",
            )
            for column in range(len(labels))
        )
        ws.append(row)


def populate_absolute_workbook(metadata: list[StateMeta], data: InputData) -> None:
    """Create the complete absolute workbook in constant-memory write-only mode."""
    workbook = Workbook(write_only=True)
    workbook.add_named_style(NamedStyle(name="absolute_header", font=Font(bold=True)))
    workbook.add_named_style(
        NamedStyle(name="absolute_scientific", number_format=SCIENTIFIC_FORMAT)
    )

    add_absolute_comment_sheet(workbook)
    add_state_values_sheet(workbook, metadata, data)
    labels = [item.label for item in metadata]
    add_absolute_covariance_sheet(
        workbook,
        "Covariance_posterior_absolute",
        labels,
        data.Sigma_f,
    )
    add_absolute_covariance_sheet(
        workbook,
        "Covariance_prior_absolute",
        labels,
        data.Sigma_i,
    )

    workbook.save(ABSOLUTE_OUTPUT_PATH)
    workbook.close()


def validate_saved_workbook(path: Path, expected_sheets: list[str]) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Expected output workbook was not created: {path}")
    workbook = load_workbook(str(path), read_only=True, data_only=False)
    try:
        missing = [name for name in expected_sheets if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"{path}: saved workbook is missing sheet(s) {missing}")
    finally:
        workbook.close()


def validate_outputs() -> None:
    """Reopen both generated files once and verify their required sheet names."""
    validate_saved_workbook(RELATIVE_OUTPUT_PATH, RELATIVE_SHEET_NAMES)
    validate_saved_workbook(ABSOLUTE_OUTPUT_PATH, ABSOLUTE_SHEET_NAMES)


def main() -> None:
    data = load_inputs()
    metadata, index_by_label = build_state_metadata()
    selection = build_template_selection(index_by_label)
    relative = calculate_relative_outputs(data, selection)
    diagnostics = validate_relative_outputs(data, selection, relative)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    populate_relative_workbook(data, relative, selection, index_by_label)
    populate_absolute_workbook(metadata, data)
    validate_outputs()

    print(f"Loaded {EXPECTED_STATE_SIZE}-state, three-response Exercise-2 result.")
    print(
        f"Template state coverage: {selection.indices.size} / {EXPECTED_STATE_SIZE}."
    )
    print("Omitted from NEA relative template: MT456 (51 values).")
    print(f"Wrote: {RELATIVE_OUTPUT_PATH}")
    print(f"Wrote: {ABSOLUTE_OUTPUT_PATH}")
    print(
        "Undefined relative mean cells written as NaN: "
        + f"{diagnostics.undefined_mean_count}"
    )
    print(
        "Undefined prior relative covariance cells written as NaN: "
        + f"{diagnostics.undefined_prior_covariance_count}"
    )
    print(
        "Undefined posterior relative covariance cells written as NaN: "
        + f"{diagnostics.undefined_posterior_covariance_count}"
    )
    print(
        "Very large finite relative values "
        + f"(|value| > {VERY_LARGE_RELATIVE_THRESHOLD:.1e}): "
        + f"{diagnostics.very_large_finite_count}"
    )
    print("Export validation passed.")


if __name__ == "__main__":
    main()
